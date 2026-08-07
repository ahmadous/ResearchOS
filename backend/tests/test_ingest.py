"""Tests d'import de fichiers — extraction (PDF/Word/Excel/texte) + pièces jointes.

On génère de vrais petits fichiers en mémoire. Lancer : python tests/test_ingest.py
"""
import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app import create_app
from app.ingest import extract, kind_of


# --- Générateurs de fichiers de test ---
def make_docx(text):
    import docx
    d = docx.Document()
    for line in text.split("\n"):
        d.add_paragraph(line)
    buf = io.BytesIO(); d.save(buf); return buf.getvalue()


def make_xlsx(rows):
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active
    for r in rows:
        ws.append(r)
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def make_pdf(text):
    import fitz
    doc = fitz.open(); page = doc.new_page()
    page.insert_text((72, 72), text)
    return doc.tobytes()


def make_png():
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (8, 8), (100, 150, 200)).save(buf, "PNG")
    return buf.getvalue()


# --- Extraction pure ---
def test_kind_detection():
    assert kind_of("a.pdf") == "pdf" and kind_of("b.docx") == "word"
    assert kind_of("c.xlsx") == "excel" and kind_of("d.png") == "image"
    assert kind_of("e.mp4") == "video" and kind_of("f.md") == "markdown"
    print("[kind]     détection par extension OK")


def test_extract_word():
    r = extract("doc.docx", make_docx("Bonjour le monde\nDeuxième ligne"))
    assert r["kind"] == "word" and "Bonjour le monde" in r["text"]
    print(f"[word]     texte extrait ({len(r['text'])} car.)")


def test_extract_excel():
    r = extract("data.xlsx", make_xlsx([["ville", "pop"], ["Dakar", 1200000]]))
    assert r["kind"] == "excel" and "Dakar" in r["text"] and "1200000" in r["text"]
    print("[excel]    cellules extraites")


def test_extract_pdf():
    r = extract("paper.pdf", make_pdf("Attention Is All You Need"))
    assert r["kind"] == "pdf" and "Attention" in r["text"]
    print(f"[pdf]      texte extrait ({len(r['text'])} car.)")


def test_image_is_binary_no_text():
    r = extract("photo.png", make_png())
    assert r["kind"] == "image" and r["text"] == ""
    print("[image]    reconnue, non extraite (pièce jointe)")


# --- Import HTTP ---
def _auth():
    c = create_app("test").test_client()
    r = c.post("/api/auth/register", json={"email": "up@u.sn", "password": "motdepasse123"})
    return c, {"Authorization": f"Bearer {r.get_json()['access_token']}"}


def test_upload_word_indexes_into_rag():
    c, h = _auth()
    r = c.post("/api/rag/upload", headers=h, content_type="multipart/form-data",
               data={"file": (io.BytesIO(make_docx("Le Transformer utilise l'attention. " * 5)), "cours.docx")})
    body = r.get_json()
    assert r.status_code == 201 and body["indexed"] is True and body["n_chunks"] >= 1
    assert body["source_type"] == "word"
    docs = c.get("/api/rag/documents", headers=h).get_json()["documents"]
    assert any(d["title"] == "cours.docx" for d in docs)
    print(f"[upload]   .docx -> indexé RAG ({body['n_chunks']} chunks)")


def test_upload_image_stored_and_served():
    c, h = _auth()
    r = c.post("/api/rag/upload", headers=h, content_type="multipart/form-data",
               data={"file": (io.BytesIO(make_png()), "photo.png")})
    body = r.get_json()
    assert r.status_code == 201 and body["indexed"] is False and body["is_attachment"]
    # Le fichier est servi
    f = c.get(f"/api/rag/documents/{body['id']}/file", headers=h)
    assert f.status_code == 200 and f.data[:4] == b"\x89PNG"
    print("[upload]   .png -> pièce jointe stockée + servie")


if __name__ == "__main__":
    for t in [test_kind_detection, test_extract_word, test_extract_excel,
              test_extract_pdf, test_image_is_binary_no_text,
              test_upload_word_indexes_into_rag, test_upload_image_stored_and_served]:
        t()
    print("\n✅ 7 tests import fichiers passés.")
